"""
Document loaders for PDF, DOCX, and text resumes.
Extracts clean, normalized text from candidate files.
"""
import io
import os
from typing import Union, BinaryIO


def extract_text_from_pdf(file_input: Union[str, BinaryIO, bytes]) -> str:
    """
    Extracts text content from a PDF file using pypdf.
    Accepts a filepath, file-like buffer, or raw bytes.
    """
    try:
        import pypdf

        if isinstance(file_input, (str, bytes)):
            stream = io.BytesIO(file_input) if isinstance(file_input, bytes) else open(file_input, "rb")
        else:
            stream = file_input

        reader = pypdf.PdfReader(stream)
        extracted_pages = []
        for index, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                extracted_pages.append(page_text.strip())

        full_text = "\n\n".join(extracted_pages).strip()
        if not full_text:
            raise ValueError("The PDF appears to be empty or contains scanned images without selectable text.")
        return full_text
    except Exception as e:
        raise RuntimeError(f"Failed to extract text from PDF: {str(e)}") from e


def extract_text_from_docx(file_input: Union[str, BinaryIO, bytes]) -> str:
    """
    Extracts text content from a DOCX file using python-docx.
    Accepts a filepath, file-like buffer, or raw bytes.
    """
    try:
        import docx

        if isinstance(file_input, bytes):
            stream = io.BytesIO(file_input)
        elif isinstance(file_input, str):
            stream = file_input
        else:
            stream = file_input

        document = docx.Document(stream)
        paragraphs = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        
        # Also extract table text if present
        for table in document.tables:
            for row in table.rows:
                row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                if row_text:
                    paragraphs.append(row_text)

        full_text = "\n\n".join(paragraphs).strip()
        if not full_text:
            raise ValueError("The DOCX document contains no readable text.")
        return full_text
    except Exception as e:
        raise RuntimeError(f"Failed to extract text from DOCX: {str(e)}") from e


def extract_text_from_file(file_obj, filename: str) -> str:
    """
    Unified dispatcher to extract text from uploaded files based on extension.
    Supported: .pdf, .docx, .txt, .md
    """
    if not filename:
        raise ValueError("Filename must be provided.")

    ext = os.path.splitext(filename)[1].lower()
    
    if ext == ".pdf":
        return extract_text_from_pdf(file_obj)
    elif ext in [".docx", ".doc"]:
        return extract_text_from_docx(file_obj)
    elif ext in [".txt", ".md"]:
        if hasattr(file_obj, "read"):
            content = file_obj.read()
            if isinstance(content, bytes):
                return content.decode("utf-8", errors="ignore").strip()
            return str(content).strip()
        elif isinstance(file_obj, str):
            with open(file_obj, "r", encoding="utf-8", errors="ignore") as f:
                return f.read().strip()
        elif isinstance(file_obj, bytes):
            return file_obj.decode("utf-8", errors="ignore").strip()
    
    raise ValueError(f"Unsupported file format '{ext}'. Please upload a PDF or DOCX file.")
